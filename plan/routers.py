from fastapi import APIRouter, HTTPException, Depends
from typing import Optional, List
from datetime import datetime
import io
import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi.responses import StreamingResponse

from database import get_db
from auth import get_current_user, hash_password, verify_password, create_access_token, validate_email_format, validate_password_strength
from models import (
    UserRegister, UserOut, UserLogin, TokenOut, ProfileUpdate,
    ProjectCreate, ProjectOut, TaskCreate, TaskOut, TaskUpdate,
    CategoryCreate, CategoryOut, TagCreate, TagOut,
    ClientCreate, ClientUpdate, ClientOut,
    ReportCreate, ReportOut, TeamCreate, TeamOut, TeamMemberOut
)

router = APIRouter(prefix="/api")

@router.post("/register", response_model=UserOut)
def register(user: UserRegister, conn=Depends(get_db)):
    if not validate_email_format(user.email):
        raise HTTPException(400, "Неверный формат email")
    if not validate_password_strength(user.password):
        raise HTTPException(400, "Пароль должен содержать минимум 6 символов")
    cur = conn.cursor()
    cur.execute("SELECT id_user FROM users WHERE email = %s", (user.email,))
    if cur.fetchone():
        raise HTTPException(400, "Email already registered")
    hashed_password = hash_password(user.password)
    cur.execute(
        "INSERT INTO users (username, email, password, user_type, notif_p) VALUES (%s, %s, %s, %s, %s) RETURNING id_user, username, email, user_type",
        (user.username, user.email, hashed_password, user.user_type, 1)
    )
    new_user = cur.fetchone()
    conn.commit()
    cur.close()
    return UserOut(id=new_user["id_user"], username=new_user["username"], email=new_user["email"], user_type=new_user["user_type"])

@router.post("/login", response_model=TokenOut)
def login(user: UserLogin, conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT id_user, password, user_type FROM users WHERE email = %s", (user.email,))
    db_user = cur.fetchone()
    cur.close()
    if not db_user or not verify_password(user.password, db_user["password"]):
        raise HTTPException(401, "Invalid email or password")
    token = create_access_token(db_user["id_user"])
    return {"access_token": token, "token_type": "bearer", "user_type": db_user["user_type"]}

@router.get("/profile")
def get_profile(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT id_user, username, email, user_type, language, notif_p FROM users WHERE id_user = %s", (current_user["id_user"],))
    user = cur.fetchone()
    cur.close()
    return {
        "id": user["id_user"],
        "username": user["username"],
        "email": user["email"],
        "user_type": user["user_type"],
        "language": user["language"],
        "notif_p": user["notif_p"]
    }

@router.put("/profile")
def update_profile(profile: ProfileUpdate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    if profile.username:
        cur.execute("UPDATE users SET username = %s WHERE id_user = %s", (profile.username, current_user["id_user"]))
    if profile.language:
        cur.execute("UPDATE users SET language = %s WHERE id_user = %s", (profile.language, current_user["id_user"]))
    if profile.notif_p is not None:
        cur.execute("UPDATE users SET notif_p = %s WHERE id_user = %s", (profile.notif_p, current_user["id_user"]))
    conn.commit()
    cur.close()
    return {"message": "Profile updated"}

@router.post("/projects", response_model=ProjectOut)
def create_project(project: ProjectCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    if project.type == "team":
        cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (project.team_id, current_user["id_user"]))
        member = cur.fetchone()
        if not member or member["role"] != "manager":
            raise HTTPException(403, "Only team managers can create team projects")
        cur.execute(
            "INSERT INTO projects (project_name, description, author_id, project_type, team_id) VALUES (%s, %s, %s, 'team', %s) RETURNING project_id",
            (project.name, project.description, current_user["id_user"], project.team_id)
        )
        new_id = cur.fetchone()["project_id"]
    else:
        cur.execute(
            "INSERT INTO projects (project_name, description, author_id, project_type) VALUES (%s, %s, %s, 'personal') RETURNING project_id",
            (project.name, project.description, current_user["id_user"])
        )
        new_id = cur.fetchone()["project_id"]
    conn.commit()
    cur.close()
    return ProjectOut(
        id=new_id, name=project.name, description=project.description,
        status="active", type=project.type, team_id=project.team_id if project.type == "team" else None,
        team_name=None, category_id=None, client_id=None
    )

@router.get("/projects", response_model=List[ProjectOut])
def get_projects(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    if current_user["user_type"] == 2:
        cur.execute("""
            SELECT p.project_id, p.project_name, p.description, 'personal' as type,
                   NULL as team_id, NULL as team_name
            FROM projects p
            WHERE p.author_id = %s
        """, (current_user["id_user"],))
    else:
        cur.execute("""
            SELECT p.project_id, p.project_name, p.description, 
                   p.project_type as type,
                   t.team_id, t.team_name
            FROM projects p
            LEFT JOIN teams t ON p.team_id = t.team_id
            LEFT JOIN team_members tm ON t.team_id = tm.team_id AND tm.user_id = %s
            WHERE p.author_id = %s OR tm.user_id IS NOT NULL
        """, (current_user["id_user"], current_user["id_user"]))
    rows = cur.fetchall()
    cur.close()
    return [ProjectOut(
        id=r["project_id"], 
        name=r["project_name"], 
        description=r["description"],
        status="active", 
        type=r["type"], 
        team_id=r["team_id"],
        team_name=r["team_name"],
        category_id=None,
        client_id=None
    ) for r in rows]

@router.delete("/projects/{project_id}")
def delete_project(project_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT author_id FROM projects WHERE project_id = %s", (project_id,))
    project = cur.fetchone()
    if not project:
        raise HTTPException(404, "Project not found")
    if project["author_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your project")
    cur.execute("DELETE FROM tasks WHERE task_project_id = %s", (project_id,))
    cur.execute("DELETE FROM projects WHERE project_id = %s", (project_id,))
    conn.commit()
    cur.close()
    return {"message": "Project deleted"}

@router.post("/tasks", response_model=TaskOut)
def create_task(task: TaskCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    project_id = task.project_id if task.project_id else None
    due_date = task.due_date if task.due_date else datetime.now().strftime("%Y-%m-%d")
    task_time = None
    if task.task_time and task.task_time.strip():
        task_time = task.task_time.strip()[:5]
    
    if project_id:
        cur.execute("SELECT team_id FROM projects WHERE project_id = %s", (project_id,))
        proj = cur.fetchone()
        if proj and proj["team_id"]:
            if not task.assigned_to:
                raise HTTPException(400, "Team tasks must have an assigned user")
            cur.execute("SELECT 1 FROM team_members WHERE team_id = %s AND user_id = %s", (proj["team_id"], task.assigned_to))
            if not cur.fetchone():
                raise HTTPException(400, "Assigned user is not a member of this team")
    
    cur.execute(
        """INSERT INTO tasks (task_name, task_description, task_priority, task_status,
           task_project_id, task_autor_id, task_date, task_tag_id, task_time, assigned_to)
           VALUES (%s, %s, %s, 'active', %s, %s, %s, %s, %s, %s)
           RETURNING task_id, task_name, task_description, task_priority, task_status, 
                      task_project_id, task_date, task_tag_id, 
                      TO_CHAR(task_time, 'HH24:MI') as task_time, assigned_to""",
        (task.title, task.description, task.priority, project_id, 
         current_user["id_user"], due_date, task.tag_id, task_time, task.assigned_to)
    )
    new = cur.fetchone()
    conn.commit()
    
    tag_name = None
    if new["task_tag_id"]:
        cur.execute("SELECT tag_name FROM tags WHERE tag_id = %s", (new["task_tag_id"],))
        tag = cur.fetchone()
        if tag:
            tag_name = tag["tag_name"]
    
    assigned_to_username = None
    if new["assigned_to"]:
        cur.execute("SELECT username FROM users WHERE id_user = %s", (new["assigned_to"],))
        user = cur.fetchone()
        if user:
            assigned_to_username = user["username"]
    
    cur.close()
    return TaskOut(
        id=new["task_id"], 
        title=new["task_name"], 
        description=new["task_description"],
        priority=new["task_priority"], 
        status=new["task_status"], 
        project_id=new["task_project_id"],
        due_date=str(new["task_date"].date()) if new["task_date"] else None,
        tag_id=new["task_tag_id"], 
        tag_name=tag_name,
        task_time=new["task_time"] if new["task_time"] else None,
        assigned_to=new["assigned_to"],
        assigned_to_username=assigned_to_username
    )

@router.get("/tasks", response_model=List[TaskOut])
def get_tasks(
    current_user: dict = Depends(get_current_user),
    project_id: Optional[int] = None,
    tag_id: Optional[int] = None,
    date: Optional[str] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    priority: Optional[str] = None,
    conn=Depends(get_db)
):
    cur = conn.cursor()
    query = """SELECT t.task_id, t.task_name, t.task_description, t.task_priority, t.task_status,
                      t.task_project_id, t.task_date::date as task_date, t.task_tag_id, tg.tag_name,
                      TO_CHAR(t.task_time, 'HH24:MI') as task_time, t.assigned_to, u.username as assigned_to_username
               FROM tasks t
               LEFT JOIN tags tg ON t.task_tag_id = tg.tag_id
               LEFT JOIN projects p ON t.task_project_id = p.project_id
               LEFT JOIN teams te ON p.team_id = te.team_id
               LEFT JOIN team_members tm ON te.team_id = tm.team_id AND tm.user_id = %s
               LEFT JOIN users u ON t.assigned_to = u.id_user
               WHERE (t.task_autor_id = %s OR tm.user_id IS NOT NULL)"""
    params = [current_user["id_user"], current_user["id_user"]]
    if project_id:
        query += " AND t.task_project_id = %s"
        params.append(project_id)
    if tag_id:
        query += " AND t.task_tag_id = %s"
        params.append(tag_id)
    if date:
        query += " AND t.task_date::date = %s"
        params.append(date)
    if status:
        query += " AND t.task_status = %s"
        params.append(status)
    if priority:
        query += " AND t.task_priority = %s"
        params.append(priority)
    if search:
        query += " AND (t.task_name ILIKE %s OR t.task_description ILIKE %s)"
        params.extend([f"%{search}%", f"%{search}%"])
    query += """ GROUP BY t.task_id, t.task_name, t.task_description, t.task_priority, 
                        t.task_status, t.task_project_id, t.task_date, t.task_tag_id, 
                        tg.tag_name, t.task_time, t.assigned_to, u.username
                 ORDER BY t.task_date ASC, t.task_id DESC"""
    cur.execute(query, params)
    rows = cur.fetchall()
    cur.close()
    return [TaskOut(
        id=r["task_id"], 
        title=r["task_name"], 
        description=r["task_description"],
        priority=r["task_priority"], 
        status=r["task_status"], 
        project_id=r["task_project_id"],
        due_date=str(r["task_date"]) if r["task_date"] else None,
        tag_id=r["task_tag_id"], 
        tag_name=r["tag_name"],
        task_time=r["task_time"] if r["task_time"] else None,
        assigned_to=r["assigned_to"],
        assigned_to_username=r["assigned_to_username"]
    ) for r in rows]

@router.put("/tasks/{task_id}", response_model=TaskOut)
def update_task(task_id: int, task_update: TaskUpdate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT task_autor_id FROM tasks WHERE task_id = %s", (task_id,))
    task = cur.fetchone()
    if not task:
        raise HTTPException(404, "Task not found")
    if task["task_autor_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your task")
    update_fields = []
    params = []
    if task_update.status is not None:
        update_fields.append("task_status = %s")
        params.append(task_update.status)
    if task_update.title is not None:
        update_fields.append("task_name = %s")
        params.append(task_update.title)
    if task_update.description is not None:
        update_fields.append("task_description = %s")
        params.append(task_update.description)
    if task_update.priority is not None:
        update_fields.append("task_priority = %s")
        params.append(task_update.priority)
    if task_update.due_date is not None:
        update_fields.append("task_date = %s")
        params.append(task_update.due_date)
    if task_update.tag_id is not None:
        update_fields.append("task_tag_id = %s")
        params.append(task_update.tag_id)
    if update_fields:
        params.append(task_id)
        cur.execute(f"UPDATE tasks SET {', '.join(update_fields)} WHERE task_id = %s", params)
        conn.commit()
    cur.execute("""SELECT t.task_id, t.task_name, t.task_description, t.task_priority, t.task_status,
                          t.task_project_id, t.task_date::date as task_date, t.task_tag_id, tg.tag_name,
                          t.task_time::text as task_time
                   FROM tasks t LEFT JOIN tags tg ON t.task_tag_id = tg.tag_id
                   WHERE t.task_id = %s""", (task_id,))
    updated = cur.fetchone()
    cur.close()
    return TaskOut(
        id=updated["task_id"], 
        title=updated["task_name"], 
        description=updated["task_description"],
        priority=updated["task_priority"], 
        status=updated["task_status"], 
        project_id=updated["task_project_id"],
        due_date=str(updated["task_date"]) if updated["task_date"] else None,
        tag_id=updated["task_tag_id"], 
        tag_name=updated["tag_name"],
        task_time=updated["task_time"][:5] if updated["task_time"] else None
    )

@router.delete("/tasks/{task_id}")
def delete_task(task_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT task_autor_id FROM tasks WHERE task_id = %s", (task_id,))
    task = cur.fetchone()
    if not task:
        raise HTTPException(404, "Task not found")
    if task["task_autor_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your task")
    cur.execute("DELETE FROM tasks WHERE task_id = %s", (task_id,))
    conn.commit()
    cur.close()
    return {"message": "Task deleted"}

@router.post("/categories", response_model=CategoryOut)
def create_category(category: CategoryCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO categories (category_name, category_autor_id) VALUES (%s, %s) RETURNING category_id, category_name",
        (category.name, current_user["id_user"])
    )
    new = cur.fetchone()
    conn.commit()
    cur.close()
    return CategoryOut(id=new["category_id"], name=new["category_name"])

@router.get("/categories", response_model=List[CategoryOut])
def get_categories(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT category_id, category_name FROM categories WHERE category_autor_id = %s", (current_user["id_user"],))
    rows = cur.fetchall()
    cur.close()
    return [CategoryOut(id=r["category_id"], name=r["category_name"]) for r in rows]

@router.delete("/categories/{category_id}")
def delete_category(category_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT category_autor_id FROM categories WHERE category_id = %s", (category_id,))
    cat = cur.fetchone()
    if not cat:
        raise HTTPException(404, "Category not found")
    if cat["category_autor_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your category")
    cur.execute("DELETE FROM categories WHERE category_id = %s", (category_id,))
    conn.commit()
    cur.close()
    return {"message": "Category deleted"}

@router.post("/tags", response_model=TagOut)
def create_tag(tag: TagCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO tags (tag_name, tag_autor_id) VALUES (%s, %s) RETURNING tag_id, tag_name",
        (tag.name, current_user["id_user"])
    )
    new = cur.fetchone()
    conn.commit()
    cur.close()
    return TagOut(id=new["tag_id"], name=new["tag_name"])

@router.get("/tags", response_model=List[TagOut])
def get_tags(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT tag_id, tag_name FROM tags WHERE tag_autor_id = %s", (current_user["id_user"],))
    rows = cur.fetchall()
    cur.close()
    return [TagOut(id=r["tag_id"], name=r["tag_name"]) for r in rows]

@router.delete("/tags/{tag_id}")
def delete_tag(tag_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT tag_autor_id FROM tags WHERE tag_id = %s", (tag_id,))
    tag = cur.fetchone()
    if not tag:
        raise HTTPException(404, "Tag not found")
    if tag["tag_autor_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your tag")
    cur.execute("DELETE FROM tags WHERE tag_id = %s", (tag_id,))
    conn.commit()
    cur.close()
    return {"message": "Tag deleted"}

@router.post("/clients", response_model=ClientOut)
def create_client(client: ClientCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    if current_user["user_type"] == 1:
        raise HTTPException(403, "Only freelancers can create clients")
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO clients (client_name, client_contacts, client_additional, user_id) VALUES (%s, %s, %s, %s) RETURNING client_id, client_name, client_contacts, client_additional",
        (client.name, client.contacts, client.additional, current_user["id_user"])
    )
    new = cur.fetchone()
    conn.commit()
    cur.close()
    return ClientOut(id=new["client_id"], name=new["client_name"], contacts=new["client_contacts"], additional=new["client_additional"])

@router.get("/clients", response_model=List[ClientOut])
def get_clients(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    if current_user["user_type"] == 1:
        return []
    cur = conn.cursor()
    cur.execute("SELECT client_id, client_name, client_contacts, client_additional FROM clients WHERE user_id = %s", (current_user["id_user"],))
    rows = cur.fetchall()
    cur.close()
    return [ClientOut(id=r["client_id"], name=r["client_name"], contacts=r["client_contacts"], additional=r["client_additional"]) for r in rows]

@router.put("/clients/{client_id}", response_model=ClientOut)
def update_client(client_id: int, client: ClientUpdate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    if current_user["user_type"] == 1:
        raise HTTPException(403, "Only freelancers can update clients")
    cur = conn.cursor()
    cur.execute("SELECT user_id FROM clients WHERE client_id = %s", (client_id,))
    existing = cur.fetchone()
    if not existing:
        raise HTTPException(404, "Client not found")
    if existing["user_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your client")
    update_fields = []
    params = []
    if client.name is not None:
        update_fields.append("client_name = %s")
        params.append(client.name)
    if client.contacts is not None:
        update_fields.append("client_contacts = %s")
        params.append(client.contacts)
    if client.additional is not None:
        update_fields.append("client_additional = %s")
        params.append(client.additional)
    if update_fields:
        params.append(client_id)
        cur.execute(f"UPDATE clients SET {', '.join(update_fields)} WHERE client_id = %s", params)
        conn.commit()
    cur.execute("SELECT client_id, client_name, client_contacts, client_additional FROM clients WHERE client_id = %s", (client_id,))
    updated = cur.fetchone()
    cur.close()
    return ClientOut(
        id=updated["client_id"],
        name=updated["client_name"],
        contacts=updated["client_contacts"],
        additional=updated["client_additional"]
    )

@router.delete("/clients/{client_id}")
def delete_client(client_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    if current_user["user_type"] == 1:
        raise HTTPException(403, "Only freelancers can delete clients")
    cur = conn.cursor()
    cur.execute("SELECT user_id FROM clients WHERE client_id = %s", (client_id,))
    client = cur.fetchone()
    if not client:
        raise HTTPException(404, "Client not found")
    if client["user_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your client")
    cur.execute("DELETE FROM clients WHERE client_id = %s", (client_id,))
    conn.commit()
    cur.close()
    return {"message": "Client deleted"}

@router.get("/stats")
def get_stats(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("""
        SELECT COUNT(DISTINCT p.project_id) as cnt
        FROM projects p
        LEFT JOIN teams t ON p.team_id = t.team_id
        LEFT JOIN team_members tm ON t.team_id = tm.team_id AND tm.user_id = %s
        WHERE p.author_id = %s OR tm.user_id IS NOT NULL
    """, (current_user["id_user"], current_user["id_user"]))
    total_projects = cur.fetchone()["cnt"] or 0
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s", (current_user["id_user"],))
    total_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s AND task_status = 'completed'", (current_user["id_user"],))
    completed_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s AND task_status = 'active'", (current_user["id_user"],))
    active_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s AND task_date = CURRENT_DATE", (current_user["id_user"],))
    today_tasks = cur.fetchone()["cnt"]
    cur.close()
    return {
        "total_projects": total_projects,
        "total_tasks": total_tasks,
        "completed_tasks": completed_tasks,
        "active_tasks": active_tasks,
        "today_tasks": today_tasks
    }

@router.get("/stats/export")
def export_stats_excel(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) as cnt FROM projects WHERE author_id = %s", (current_user["id_user"],))
    total_projects = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s", (current_user["id_user"],))
    total_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s AND task_status = 'completed'", (current_user["id_user"],))
    completed_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s AND task_status = 'active'", (current_user["id_user"],))
    active_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s AND task_date = CURRENT_DATE", (current_user["id_user"],))
    today_tasks = cur.fetchone()["cnt"]
    cur.execute("SELECT task_priority, COUNT(*) as cnt FROM tasks WHERE task_autor_id = %s GROUP BY task_priority", (current_user["id_user"],))
    priorities = {r["task_priority"]: r["cnt"] for r in cur.fetchall()}
    cur.close()
    wb = openpyxl.Workbook()
    header_font = Font(name='Arial', bold=True, size=12, color='FFFFFF')
    header_fill = PatternFill(start_color='FF4081', end_color='FF4081', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='FFD0D8'),
        right=Side(style='thin', color='FFD0D8'),
        top=Side(style='thin', color='FFD0D8'),
        bottom=Side(style='thin', color='FFD0D8')
    )
    ws = wb.active
    ws.title = "Статистика"
    ws.merge_cells('A1:C1')
    ws['A1'] = "СТАТИСТИКА Focus"
    ws['A1'].font = Font(name='Arial', bold=True, size=18, color='C2185B')
    ws['A1'].alignment = Alignment(horizontal='center')
    headers = ["Показатель", "Значение", "%"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    data = [
        ("Проектов", total_projects),
        ("Всего задач", total_tasks),
        ("Активных", active_tasks),
        ("Выполнено", completed_tasks),
        ("На сегодня", today_tasks)
    ]
    for i, (label, val) in enumerate(data, 4):
        ws.cell(row=i, column=1, value=label).border = thin_border
        ws.cell(row=i, column=2, value=val).border = thin_border
        pct = round(val / total_tasks * 100, 1) if total_tasks > 0 else 0
        ws.cell(row=i, column=3, value=f"{pct}%").border = thin_border
    chart1 = PieChart()
    chart1.title = "Статусы задач"
    chart1.width = 15
    chart1.height = 10
    ws.cell(row=11, column=1, value="Статус").font = Font(bold=True)
    ws.cell(row=11, column=2, value="Кол-во").font = Font(bold=True)
    ws.cell(row=12, column=1, value="Выполнено")
    ws.cell(row=12, column=2, value=completed_tasks)
    ws.cell(row=13, column=1, value="Активные")
    ws.cell(row=13, column=2, value=active_tasks)
    data_ref = Reference(ws, min_col=2, min_row=11, max_row=13)
    cats_ref = Reference(ws, min_col=1, min_row=12, max_row=13)
    chart1.add_data(data_ref, titles_from_data=True)
    chart1.set_categories(cats_ref)
    ws.add_chart(chart1, "E3")
    chart2 = BarChart()
    chart2.title = "Приоритеты"
    chart2.width = 15
    chart2.height = 10
    ws.cell(row=16, column=1, value="Приоритет").font = Font(bold=True)
    ws.cell(row=16, column=2, value="Кол-во").font = Font(bold=True)
    priority_map = {"high": "Высокий", "medium": "Средний", "low": "Низкий"}
    row = 17
    for key, label in priority_map.items():
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=priorities.get(key, 0))
        row += 1
    data_ref2 = Reference(ws, min_col=2, min_row=16, max_row=19)
    cats_ref2 = Reference(ws, min_col=1, min_row=17, max_row=19)
    chart2.add_data(data_ref2, titles_from_data=True)
    chart2.set_categories(cats_ref2)
    ws.add_chart(chart2, "E18")
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=statistics.xlsx"}
    )

@router.post("/reports", response_model=ReportOut)
def create_report(report: ReportCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    if not report.report_name or not report.report_name.strip():
        raise HTTPException(400, "Report name is required")
    if not report.period_start or not report.period_end:
        raise HTTPException(400, "Period start and end dates are required")
    if report.report_type not in ["project", "client", "team"]:
        raise HTTPException(400, "Report type must be 'project', 'client', or 'team'")
    if not report.target_id:
        raise HTTPException(400, "Target ID is required")
    
    cur = conn.cursor()
    
    if current_user["user_type"] == 1:
        if report.report_type == "project":
            cur.execute("""
                SELECT p.team_id FROM projects p
                WHERE p.project_id = %s AND p.team_id IS NOT NULL
            """, (report.target_id,))
            proj = cur.fetchone()
            if not proj:
                raise HTTPException(400, "Regular users can only create reports for team projects")
            cur.execute("SELECT 1 FROM team_members WHERE team_id = %s AND user_id = %s", (proj["team_id"], current_user["id_user"]))
            if not cur.fetchone():
                raise HTTPException(403, "You are not a member of this team")
        elif report.report_type == "team":
            cur.execute("SELECT 1 FROM team_members WHERE team_id = %s AND user_id = %s", (report.target_id, current_user["id_user"]))
            if not cur.fetchone():
                raise HTTPException(403, "You are not a member of this team")
        else:
            raise HTTPException(400, "Regular users can only create project or team reports")
    elif current_user["user_type"] == 2:
        if report.report_type == "project":
            cur.execute("SELECT client_id FROM projects WHERE project_id = %s AND author_id = %s", (report.target_id, current_user["id_user"]))
            proj = cur.fetchone()
            if not proj or not proj["client_id"]:
                raise HTTPException(400, "Freelancers can only create reports for projects linked to a client")
        elif report.report_type == "client":
            cur.execute("SELECT 1 FROM clients WHERE client_id = %s AND user_id = %s", (report.target_id, current_user["id_user"]))
            if not cur.fetchone():
                raise HTTPException(403, "This client does not belong to you")
        else:
            raise HTTPException(400, "Freelancers can only create project or client reports")
    
    cur.execute(
        """INSERT INTO reports 
           (report_name, report_period_start, report_period_end, report_type, report_target_id, user_id) 
           VALUES (%s, %s, %s, %s, %s, %s) 
           RETURNING report_id, report_name, report_period_start, report_period_end, report_type, report_target_id, created_at""",
        (report.report_name.strip(), report.period_start, report.period_end, 
         report.report_type, report.target_id, current_user["id_user"])
    )
    new = cur.fetchone()
    conn.commit()
    cur.close()
    return ReportOut(
        id=new["report_id"], 
        report_name=new["report_name"],
        period_start=str(new["report_period_start"]), 
        period_end=str(new["report_period_end"]),
        report_type=new["report_type"], 
        target_id=new["report_target_id"],
        created_at=str(new["created_at"])
    )

@router.get("/reports", response_model=List[ReportOut])
def get_reports(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(""" 
        SELECT report_id, report_name, report_period_start, report_period_end, 
               report_type, report_target_id, created_at 
        FROM reports 
        WHERE user_id = %s 
        ORDER BY created_at DESC
    """, (current_user["id_user"],))
    rows = cur.fetchall()
    cur.close()
    return [ReportOut(
        id=r["report_id"], report_name=r["report_name"],
        period_start=str(r["report_period_start"]), period_end=str(r["report_period_end"]),
        report_type=r["report_type"], target_id=r["report_target_id"],
        created_at=str(r["created_at"])
    ) for r in rows]

@router.delete("/reports/{report_id}")
def delete_report(report_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
   
    cur = conn.cursor()
    cur.execute("SELECT user_id FROM reports WHERE report_id = %s", (report_id,))
    rep = cur.fetchone()
    if not rep:
        raise HTTPException(404, "Report not found")
    if rep["user_id"] != current_user["id_user"]:
        raise HTTPException(403, "Not your report")
    cur.execute("DELETE FROM reports WHERE report_id = %s", (report_id,))
    conn.commit()
    cur.close()
    return {"message": "Report deleted"}

@router.get("/reports/{report_id}/download")
def download_report(report_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT * FROM reports WHERE report_id = %s AND user_id = %s", (report_id, current_user["id_user"]))
    report = cur.fetchone()
    if not report:
        raise HTTPException(404, "Report not found")
    
    if report["report_type"] == "project":
        cur.execute("""
            SELECT t.task_name, t.task_description, t.task_priority, t.task_status, 
                   t.task_date::date as task_date, tg.tag_name, p.project_name
            FROM tasks t
            LEFT JOIN tags tg ON t.task_tag_id = tg.tag_id
            LEFT JOIN projects p ON t.task_project_id = p.project_id
            WHERE t.task_project_id = %s 
            AND t.task_date >= %s 
            AND t.task_date <= %s
            AND t.task_autor_id = %s
            ORDER BY t.task_date DESC
        """, (report["report_target_id"], report["report_period_start"], report["report_period_end"], current_user["id_user"]))
    else:
        cur.execute("""
            SELECT t.task_name, t.task_description, t.task_priority, t.task_status, 
                   t.task_date::date as task_date, tg.tag_name, p.project_name
            FROM tasks t
            LEFT JOIN tags tg ON t.task_tag_id = tg.tag_id
            LEFT JOIN projects p ON t.task_project_id = p.project_id
            WHERE p.client_id = %s 
            AND t.task_date >= %s 
            AND t.task_date <= %s
            AND t.task_autor_id = %s
            ORDER BY t.task_date DESC
        """, (report["report_target_id"], report["report_period_start"], report["report_period_end"], current_user["id_user"]))
    tasks_data = cur.fetchall()
    cur.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчет"
    
    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='1d1d1f', end_color='1d1d1f', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    cell_alignment = Alignment(vertical='center', wrap_text=True)
    border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    total = len(tasks_data)
    completed = sum(1 for t in tasks_data if t["task_status"] == "completed")
    active = sum(1 for t in tasks_data if t["task_status"] == "active")
    high = sum(1 for t in tasks_data if t["task_priority"] == "high")
    medium = sum(1 for t in tasks_data if t["task_priority"] == "medium")
    low = sum(1 for t in tasks_data if t["task_priority"] == "low")
    
    row = 1
    
    ws.merge_cells('A1:H1')
    ws['A1'] = report["report_name"]
    ws['A1'].font = Font(name='Arial', bold=True, size=16, color='1d1d1f')
    ws['A1'].alignment = Alignment(horizontal='center')
    row += 1
    
    ws.merge_cells('A2:H2')
    ws['A2'] = f"Период: {report['report_period_start']} — {report['report_period_end']}"
    ws['A2'].alignment = Alignment(horizontal='center')
    row += 1
    
    ws.merge_cells('A3:H3')
    type_text = "По проекту" if report["report_type"] == "project" else "По клиенту"
    ws['A3'] = f"Тип: {type_text}"
    ws['A3'].alignment = Alignment(horizontal='center')
    row += 1
    
    ws.merge_cells('A4:H4')
    if report["report_type"] == "project":
        cur2 = conn.cursor()
        cur2.execute("SELECT project_name FROM projects WHERE project_id = %s", (report["report_target_id"],))
        proj = cur2.fetchone()
        cur2.close()
        ws['A4'] = f"Проект: {proj['project_name'] if proj else '—'}"
    else:
        cur2 = conn.cursor()
        cur2.execute("SELECT client_name FROM clients WHERE client_id = %s", (report["report_target_id"],))
        cl = cur2.fetchone()
        cur2.close()
        ws['A4'] = f"Клиент: {cl['client_name'] if cl else '—'}"
    ws['A4'].alignment = Alignment(horizontal='center')
    row += 2
    
    headers = ["Название", "Проект", "Приоритет", "Статус", "Дата", "Тег", "Описание"]
    col_widths = [25, 20, 12, 14, 14, 14, 35]
    
    for col, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    row += 1
    
    priority_map = {"low": "Низкий", "medium": "Средний", "high": "Высокий"}
    status_map = {"active": "Активна", "completed": "Выполнена"}
    
    for task in tasks_data:
        data = [
            task["task_name"],
            task.get("project_name", "—"),
            priority_map.get(task["task_priority"], task["task_priority"]),
            status_map.get(task["task_status"], task["task_status"]),
            str(task["task_date"]) if task["task_date"] else "—",
            task["tag_name"] or "—",
            task["task_description"] or "—"
        ]
        for col_idx, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.font = Font(name='Arial', size=10)
            cell.alignment = cell_alignment
            cell.border = border
        row += 1
    
    row += 2
    
    ws.merge_cells('J1:L1')
    ws['J1'] = "СТАТИСТИКА"
    ws['J1'].font = Font(name='Arial', bold=True, size=14, color='1d1d1f')
    ws['J1'].alignment = Alignment(horizontal='center')
    
    stat_headers = ["Показатель", "Значение", "%"]
    stat_row = 3
    for col, header in enumerate(stat_headers, 10):
        cell = ws.cell(row=stat_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    stat_data = [
        ("Всего задач", total),
        ("Выполнено", completed),
        ("Активных", active),
        ("Высокий приоритет", high),
        ("Средний приоритет", medium),
        ("Низкий приоритет", low)
    ]
    
    stat_row += 1
    for label, value in stat_data:
        cell_label = ws.cell(row=stat_row, column=10, value=label)
        cell_label.font = Font(name='Arial', size=10)
        cell_label.border = border
        
        cell_value = ws.cell(row=stat_row, column=11, value=value)
        cell_value.font = Font(name='Arial', size=10)
        cell_value.alignment = Alignment(horizontal='center')
        cell_value.border = border
        
        percent = round((value / total * 100), 1) if total > 0 else 0
        cell_pct = ws.cell(row=stat_row, column=12, value=f"{percent}%")
        cell_pct.font = Font(name='Arial', size=10)
        cell_pct.alignment = Alignment(horizontal='center')
        cell_pct.border = border
        stat_row += 1
    
    if completed > 0 or active > 0:
        ws.merge_cells(f'A{row}:H{row}')
        ws[f'A{row}'] = "СТАТУСЫ ЗАДАЧ"
        ws[f'A{row}'].font = Font(name='Arial', bold=True, size=12)
        ws[f'A{row}'].fill = PatternFill(start_color='F5F5F7', end_color='F5F5F7', fill_type='solid')
        row += 2
        
        ws.cell(row=row, column=1, value="Статус").font = Font(bold=True)
        ws.cell(row=row, column=2, value="Кол-во").font = Font(bold=True)
        
        data_row = row + 1
        if completed > 0:
            ws.cell(row=data_row, column=1, value="Выполнено")
            ws.cell(row=data_row, column=2, value=completed)
            data_row += 1
        if active > 0:
            ws.cell(row=data_row, column=1, value="Активные")
            ws.cell(row=data_row, column=2, value=active)
            data_row += 1
        
        from openpyxl.chart import PieChart, Reference
        from openpyxl.chart.series import DataPoint
        
        chart = PieChart()
        chart.title = "Статусы задач"
        chart.width = 12
        chart.height = 9
        
        data_ref = Reference(ws, min_col=2, min_row=row, max_row=data_row-1)
        cats_ref = Reference(ws, min_col=1, min_row=row+1, max_row=data_row-1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        
        if len(chart.series) > 0:
            series = chart.series[0]
            series.graphicalProperties.solidFill = None
            
            idx = 0
            if completed > 0:
                point = DataPoint(idx=idx)
                point.graphicalProperties.solidFill = "16a34a"
                series.data_points.append(point)
                idx += 1
            if active > 0:
                point = DataPoint(idx=idx)
                point.graphicalProperties.solidFill = "e91e63"
                series.data_points.append(point)
        
        ws.add_chart(chart, f"D{row}")
    
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 35
    ws.column_dimensions['J'].width = 20
    ws.column_dimensions['K'].width = 12
    ws.column_dimensions['L'].width = 12
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"report_{report_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.post("/teams", response_model=TeamOut)
def create_team(team: TeamCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO teams (team_name, description, created_by) VALUES (%s, %s, %s) RETURNING team_id, team_name, description, created_by, created_at",
        (team.team_name, team.description, current_user["id_user"])
    )
    new_team = cur.fetchone()
    cur.execute(
        "INSERT INTO team_members (team_id, user_id, role) VALUES (%s, %s, 'manager')",
        (new_team["team_id"], current_user["id_user"])
    )
    conn.commit()
    cur.close()
    return TeamOut(
        team_id=new_team["team_id"],
        team_name=new_team["team_name"],
        description=new_team["description"],
        created_by=new_team["created_by"],
        created_at=str(new_team["created_at"])
    )

@router.get("/teams", response_model=List[TeamOut])
def get_teams(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("""
        SELECT t.team_id, t.team_name, t.description, t.created_by, t.created_at
        FROM teams t
        JOIN team_members tm ON t.team_id = tm.team_id
        WHERE tm.user_id = %s
        ORDER BY t.created_at DESC
    """, (current_user["id_user"],))
    rows = cur.fetchall()
    cur.close()
    return [TeamOut(
        team_id=r["team_id"],
        team_name=r["team_name"],
        description=r["description"],
        created_by=r["created_by"],
        created_at=str(r["created_at"])
    ) for r in rows]

@router.get("/teams/{team_id}", response_model=TeamOut)
def get_team(team_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT * FROM teams WHERE team_id = %s", (team_id,))
    team = cur.fetchone()
    if not team:
        raise HTTPException(404, "Team not found")
    cur.execute("SELECT 1 FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    if not cur.fetchone():
        raise HTTPException(403, "You are not a member of this team")
    cur.close()
    return TeamOut(
        team_id=team["team_id"],
        team_name=team["team_name"],
        description=team["description"],
        created_by=team["created_by"],
        created_at=str(team["created_at"])
    )

@router.put("/teams/{team_id}", response_model=TeamOut)
def update_team(team_id: int, team: TeamCreate, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT created_by FROM teams WHERE team_id = %s", (team_id,))
    existing = cur.fetchone()
    if not existing:
        raise HTTPException(404, "Team not found")
    if existing["created_by"] != current_user["id_user"]:
        raise HTTPException(403, "Only the creator can update the team")
    cur.execute(
        "UPDATE teams SET team_name = %s, description = %s WHERE team_id = %s RETURNING *",
        (team.team_name, team.description, team_id)
    )
    updated = cur.fetchone()
    conn.commit()
    cur.close()
    return TeamOut(
        team_id=updated["team_id"],
        team_name=updated["team_name"],
        description=updated["description"],
        created_by=updated["created_by"],
        created_at=str(updated["created_at"])
    )

@router.delete("/teams/{team_id}")
def delete_team(team_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT created_by FROM teams WHERE team_id = %s", (team_id,))
    team = cur.fetchone()
    if not team:
        raise HTTPException(404, "Team not found")
    if team["created_by"] != current_user["id_user"]:
        raise HTTPException(403, "Only the creator can delete the team")
    cur.execute("UPDATE projects SET team_id = NULL WHERE team_id = %s", (team_id,))
    cur.execute("DELETE FROM team_members WHERE team_id = %s", (team_id,))
    cur.execute("DELETE FROM notifications WHERE related_team_id = %s", (team_id,))
    cur.execute("DELETE FROM teams WHERE team_id = %s", (team_id,))
    conn.commit()
    cur.close()
    return {"message": "Team deleted"}

@router.get("/teams/{team_id}/members", response_model=List[TeamMemberOut])
def get_team_members(team_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("""
        SELECT tm.member_id, tm.user_id, u.username, tm.role, tm.joined_at
        FROM team_members tm
        JOIN users u ON tm.user_id = u.id_user
        WHERE tm.team_id = %s
        ORDER BY tm.role DESC, tm.joined_at
    """, (team_id,))
    rows = cur.fetchall()
    cur.close()
    return [TeamMemberOut(
        member_id=r["member_id"],
        user_id=r["user_id"],
        username=r["username"],
        role=r["role"],
        joined_at=str(r["joined_at"])
    ) for r in rows]

@router.get("/teams/{team_id}/projects")
def get_team_projects(team_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT project_id, project_name, description FROM projects WHERE team_id = %s", (team_id,))
    rows = cur.fetchall()
    cur.close()
    return [{"project_id": r["project_id"], "project_name": r["project_name"], "description": r["description"]} for r in rows]

@router.post("/teams/{team_id}/add-existing-project")
def add_existing_project_to_team(team_id: int, project_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    member = cur.fetchone()
    if not member or member["role"] != "manager":
        raise HTTPException(403, "Only managers can add projects to team")
    cur.execute("SELECT project_id, project_type FROM projects WHERE project_id = %s AND author_id = %s", (project_id, current_user["id_user"]))
    project = cur.fetchone()
    if not project:
        raise HTTPException(404, "Project not found or not yours")
    cur.execute("SELECT 1 FROM projects WHERE project_id = %s AND team_id = %s", (project_id, team_id))
    if cur.fetchone():
        raise HTTPException(400, "Project already in team")
    cur.execute("UPDATE projects SET team_id = %s, project_type = 'team' WHERE project_id = %s", (team_id, project_id))
    conn.commit()
    cur.close()
    return {"message": "Project added to team and became shared"}

@router.post("/teams/{team_id}/projects")
def add_team_project(team_id: int, project_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    member = cur.fetchone()
    if not member or member["role"] != "manager":
        raise HTTPException(403, "Only managers can add projects to team")
    cur.execute("SELECT author_id FROM projects WHERE project_id = %s", (project_id,))
    proj = cur.fetchone()
    if not proj:
        raise HTTPException(404, "Project not found")
    if proj["author_id"] != current_user["id_user"]:
        raise HTTPException(403, "You can only add your own projects")
    cur.execute("UPDATE projects SET team_id = %s, project_type = 'team' WHERE project_id = %s", (team_id, project_id))
    conn.commit()
    cur.execute("SELECT project_name, description FROM projects WHERE project_id = %s", (project_id,))
    p = cur.fetchone()
    cur.close()
    return {"project_id": project_id, "project_name": p["project_name"], "description": p["description"]}

@router.get("/teams/{team_id}/pending-invites")
def get_pending_invites(team_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("""
        SELECT n.notification_id, n.user_id, u.username, u.email
        FROM notifications n
        JOIN users u ON n.user_id = u.id_user
        WHERE n.related_team_id = %s AND n.type = 'invite' AND n.is_read = FALSE
    """, (team_id,))
    rows = cur.fetchall()
    cur.close()
    return [{"id": r["notification_id"], "user_id": r["user_id"], "username": r["username"], "email": r["email"]} for r in rows]

@router.post("/teams/{team_id}/members", response_model=TeamMemberOut)
def add_team_member(team_id: int, user_email: str, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    member = cur.fetchone()
    if not member or member["role"] != "manager":
        raise HTTPException(403, "Only managers can add members")
    cur.execute("SELECT id_user, username FROM users WHERE email = %s", (user_email,))
    user_to_add = cur.fetchone()
    if not user_to_add:
        raise HTTPException(404, "User not found")
    cur.execute(
        "INSERT INTO team_members (team_id, user_id, role) VALUES (%s, %s, 'member') RETURNING member_id, user_id, role, joined_at",
        (team_id, user_to_add["id_user"])
    )
    new_member = cur.fetchone()
    cur.execute("DELETE FROM notifications WHERE related_team_id = %s AND user_id = %s AND type = 'invite'", 
                (team_id, user_to_add["id_user"]))
    conn.commit()
    cur.close()
    return TeamMemberOut(
        member_id=new_member["member_id"],
        user_id=new_member["user_id"],
        username=user_to_add["username"],
        role=new_member["role"],
        joined_at=str(new_member["joined_at"])
    )

@router.put("/teams/{team_id}/members/{user_id}/role")
def update_member_role(team_id: int, user_id: int, role: str, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    current_member = cur.fetchone()
    if not current_member or current_member["role"] != "manager":
        raise HTTPException(403, "Only managers can change roles")
    cur.execute("SELECT created_by FROM teams WHERE team_id = %s", (team_id,))
    team = cur.fetchone()
    if team["created_by"] == user_id:
        raise HTTPException(403, "Cannot change team creator's role")
    if role not in ["manager", "member"]:
        raise HTTPException(400, "Invalid role")
    cur.execute("UPDATE team_members SET role = %s WHERE team_id = %s AND user_id = %s", (role, team_id, user_id))
    conn.commit()
    cur.close()
    return {"message": "Role updated"}

@router.delete("/teams/{team_id}/members/{user_id}")
def remove_team_member(team_id: int, user_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    member = cur.fetchone()
    if not member:
        raise HTTPException(403, "You are not a member of this team")
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, user_id))
    target = cur.fetchone()
    if not target:
        raise HTTPException(404, "Member not found")
    if user_id == current_user["id_user"]:
        raise HTTPException(403, "You cannot remove yourself from the team")
    if member["role"] != "manager" and target["role"] == "manager":
        raise HTTPException(403, "Only managers can remove members")
    cur.execute("DELETE FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, user_id))
    conn.commit()
    cur.close()
    return {"message": "Member removed"}

@router.post("/teams/{team_id}/invite")
def invite_user(team_id: int, user_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT role FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, current_user["id_user"]))
    member = cur.fetchone()
    if not member or member["role"] != "manager":
        raise HTTPException(403, "Only managers can invite")
    cur.execute("SELECT 1 FROM team_members WHERE team_id = %s AND user_id = %s", (team_id, user_id))
    if cur.fetchone():
        raise HTTPException(400, "User already in team")
    cur.execute("SELECT user_type, notif_p FROM users WHERE id_user = %s", (user_id,))
    invited_user = cur.fetchone()
    if invited_user and invited_user["user_type"] == 2:
        raise HTTPException(400, "Freelancer cannot be invited to teams")
    cur.execute("SELECT team_name FROM teams WHERE team_id = %s", (team_id,))
    team = cur.fetchone()
    team_name = team["team_name"] if team else f"команда #{team_id}"
    if invited_user and invited_user["notif_p"] == 1:
        cur.execute(
            "INSERT INTO notifications (user_id, message, type, related_team_id) VALUES (%s, %s, 'invite', %s)",
            (user_id, f"Вас пригласили в команду «{team_name}»", team_id)
        )
    conn.commit()
    cur.close()
    return {"message": "Invitation sent"}

@router.get("/tasks/deadline-notifications")
def get_deadline_notifications(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT notif_p FROM users WHERE id_user = %s", (current_user["id_user"],))
    user = cur.fetchone()
    if not user or user["notif_p"] != 1:
        cur.close()
        return {"notifications_created": 0}
    cur.execute("""
        SELECT t.task_id, t.task_name, TO_CHAR(t.task_time, 'HH24:MI') as task_time
        FROM tasks t
        WHERE t.task_autor_id = %s 
          AND t.task_date::date = CURRENT_DATE
          AND t.task_time IS NOT NULL
          AND t.task_time::time BETWEEN (CURRENT_TIME - INTERVAL '2 minutes') AND CURRENT_TIME
          AND t.task_status = 'active'
    """, (current_user["id_user"],))
    tasks_today = cur.fetchall()
    notifications_created = 0
    for task in tasks_today:
        cur.execute("""
            SELECT notification_id FROM notifications 
            WHERE user_id = %s AND task_id = %s AND type = 'deadline'
        """, (current_user["id_user"], task["task_id"]))
        if not cur.fetchone():
            time_val = task["task_time"] or "00:00"
            message = f"⏰ В {time_val} задача «{task['task_name']}»"
            cur.execute("""
                INSERT INTO notifications (user_id, message, type, task_id, is_read)
                VALUES (%s, %s, 'deadline', %s, FALSE)
            """, (current_user["id_user"], message, task["task_id"]))
            notifications_created += 1
    conn.commit()
    cur.close()
    return {"notifications_created": notifications_created}

@router.get("/notifications")
def get_notifications(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("""
        SELECT n.*, t.team_name
        FROM notifications n
        LEFT JOIN teams t ON n.related_team_id = t.team_id
        WHERE n.user_id = %s 
        ORDER BY n.created_at DESC
    """, (current_user["id_user"],))
    rows = cur.fetchall()
    cur.close()
    return [{
        "id": r["notification_id"], 
        "message": r["message"], 
        "is_read": r["is_read"], 
        "type": r["type"], 
        "team_id": r["related_team_id"],
        "team_name": r.get("team_name"),
        "task_id": r.get("task_id"),
        "created_at": str(r["created_at"])
    } for r in rows]

@router.put("/notifications/{notification_id}/read")
def mark_notification_read(notification_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(
        "UPDATE notifications SET is_read = TRUE WHERE notification_id = %s AND user_id = %s",
        (notification_id, current_user["id_user"])
    )
    conn.commit()
    cur.close()
    return {"message": "Marked as read"}

@router.put("/notifications/read-all")
def mark_all_notifications_read(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("UPDATE notifications SET is_read = TRUE WHERE user_id = %s AND is_read = FALSE", (current_user["id_user"],))
    conn.commit()
    cur.close()
    return {"message": "All notifications marked as read"}

@router.delete("/notifications/{notification_id}/delete")
def delete_notification(notification_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(
        "DELETE FROM notifications WHERE notification_id = %s AND user_id = %s",
        (notification_id, current_user["id_user"])
    )
    conn.commit()
    cur.close()
    return {"message": "Deleted"}

@router.delete("/notifications/delete-all")
def delete_all_notifications(current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute(
        "DELETE FROM notifications WHERE user_id = %s",
        (current_user["id_user"],)
    )
    conn.commit()
    cur.close()
    return {"message": "All notifications deleted"}

@router.post("/notifications/{notification_id}/accept")
def accept_invite(notification_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT * FROM notifications WHERE notification_id = %s AND user_id = %s", (notification_id, current_user["id_user"]))
    notif = cur.fetchone()
    if not notif:
        raise HTTPException(404, "Notification not found")
    if notif["type"] == "invite" and notif["related_team_id"]:
        cur.execute("INSERT INTO team_members (team_id, user_id, role) VALUES (%s, %s, 'member') ON CONFLICT DO NOTHING", 
                    (notif["related_team_id"], current_user["id_user"]))
    cur.execute("UPDATE notifications SET is_read = TRUE WHERE notification_id = %s", (notification_id,))
    conn.commit()
    cur.close()
    return {"message": "Accepted"}

@router.post("/api/notifications/{notification_id}/reject")
def reject_invite(notification_id: int, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("SELECT * FROM notifications WHERE notification_id = %s AND user_id = %s", (notification_id, current_user["id_user"]))
    notif = cur.fetchone()
    if not notif:
        raise HTTPException(404, "Notification not found")
    cur.execute("DELETE FROM notifications WHERE notification_id = %s", (notification_id,))
    conn.commit()
    cur.close()
    return {"message": "Invitation rejected"}

@router.get("/users/search")
def search_users(q: str, current_user: dict = Depends(get_current_user), conn=Depends(get_db)):
    cur = conn.cursor()
    cur.execute("""
        SELECT id_user, username, email 
        FROM users 
        WHERE (email ILIKE %s OR username ILIKE %s) 
        AND id_user != %s 
        LIMIT 10
    """, (f"%{q}%", f"%{q}%", current_user["id_user"]))
    rows = cur.fetchall()
    cur.close()
    return [{"id": r["id_user"], "username": r["username"], "email": r["email"]} for r in rows]